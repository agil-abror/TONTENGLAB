"""
TongTengLab Pustaka — pembuat daftar pustaka otomatis.
Jalankan:  streamlit run app.py
"""
import io
import json
import re
from datetime import date

import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from formatter import (LABELS, MONTHS, STYLES, TYPES, check_entry, clean_doi,
                       format_list, parse_authors)

st.set_page_config(page_title="TongTengLab Pustaka", page_icon="📚", layout="wide")

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:wght@500;700;800&family=IBM+Plex+Sans:wght@400;500&display=swap');
html, body, [class*="css"] { font-family: 'IBM Plex Sans', sans-serif; }
h1, h2, h3 { font-family: 'Bricolage Grotesque', sans-serif !important; letter-spacing: -0.02em; }
.block-container { padding-top: 2rem; max-width: 1100px; }
.tt-hero {
  border: 2px solid #14262B; border-radius: 14px; padding: 1.3rem 1.6rem; margin-bottom: 1.2rem;
  background:
    linear-gradient(#14262B14 1px, transparent 1px) 0 0 / 100% 22px,
    linear-gradient(90deg, #14262B14 1px, transparent 1px) 0 0 / 22px 100%, #FFFFFF;
}
.tt-hero h1 { margin: 0; font-size: 2.4rem; font-weight: 800; }
.tt-hero p { margin: .3rem 0 0; font-size: 1.05rem; max-width: 60ch; }
</style>
<div class="tt-hero">
  <h1>TongTengLab Pustaka</h1>
  <p>Isi data sumber, pilih gaya sitasi, dan daftar pustaka tersusun rapi. Bisa diunduh ke Word.</p>
</div>
""",
    unsafe_allow_html=True,
)

st.session_state.setdefault("entries", [])
st.session_state.setdefault("lk_msg", "")

with st.sidebar:
    st.header("Pengaturan")
    style = st.selectbox("Gaya sitasi", STYLES)
    lang_name = st.radio("Bahasa label", ["Indonesia", "English"], help="Memengaruhi kata seperti dan/and, hlm./pp., Diakses/Accessed, dan nama bulan.")
    lang = "id" if lang_name == "Indonesia" else "en"
    st.caption(
        "APA, MLA, dan Harvard diurutkan abjad. IEEE diurutkan sesuai urutan kamu menambahkan sumber "
        "(sesuai urutan kutipan di tulisan). Cek pedoman kampus atau dosenmu bila ada aturan khusus."
    )

TYPE_KEYS = list(TYPES.keys())


# ───────── Isi otomatis dari DOI / ISBN ─────────
def _strip_tags(x):
    return re.sub(r"<[^>]+>", "", x or "").strip()


def do_lookup():
    q = st.session_state.get("lk_id", "").strip()
    st.session_state.lk_msg = ""
    if not q:
        st.session_state.lk_msg = "Isi DOI atau ISBN dulu."
        return
    try:
        if "10." in q:
            d = clean_doi(q)
            r = requests.get(f"https://api.crossref.org/works/{d}", timeout=15,
                             headers={"User-Agent": "TongTengLabPustaka/1.0"})
            r.raise_for_status()
            m = r.json()["message"]
            authors = []
            for a in m.get("author", []):
                if a.get("family"):
                    authors.append(f"{a['family']}, {a.get('given', '')}".strip(", "))
                elif a.get("name"):
                    authors.append(a["name"])
            ptype = m.get("type", "")
            kind = "buku" if ptype in ("book", "monograph", "edited-book", "reference-book") else \
                   "skripsi" if ptype == "dissertation" else "jurnal"
            st.session_state.f_type = TYPES[kind]
            st.session_state.f_authors = "\n".join(authors)
            st.session_state.f_org = False
            year = (m.get("issued", {}).get("date-parts") or [[None]])[0][0]
            st.session_state.f_year = str(year or "")
            st.session_state.f_title = _strip_tags((m.get("title") or [""])[0])
            if kind == "jurnal":
                st.session_state.f_journal = _strip_tags((m.get("container-title") or [""])[0])
                st.session_state.f_volume = m.get("volume", "")
                st.session_state.f_issue = m.get("issue", "")
                st.session_state.f_pages = m.get("page", "")
                st.session_state.f_doi = d
            elif kind == "buku":
                st.session_state.f_publisher = m.get("publisher", "")
                st.session_state.f_doi = d
            st.session_state.lk_msg = "Data ditemukan di Crossref. Periksa kolom di bawah lalu klik Tambah ke daftar."
        else:
            isbn = re.sub(r"[^0-9Xx]", "", q)
            if len(isbn) not in (10, 13):
                st.session_state.lk_msg = "Format tidak dikenali. Gunakan DOI (mengandung 10.xxxx) atau ISBN 10/13 digit."
                return
            r = requests.get("https://openlibrary.org/api/books", timeout=15,
                             params={"bibkeys": f"ISBN:{isbn}", "format": "json", "jscmd": "data"})
            r.raise_for_status()
            data = r.json().get(f"ISBN:{isbn}")
            if not data:
                st.session_state.lk_msg = "ISBN tidak ditemukan di Open Library. Isi manual saja."
                return
            st.session_state.f_type = TYPES["buku"]
            st.session_state.f_authors = "\n".join(a["name"] for a in data.get("authors", []))
            st.session_state.f_org = False
            yr = re.search(r"\d{4}", data.get("publish_date", ""))
            st.session_state.f_year = yr.group(0) if yr else ""
            st.session_state.f_title = data.get("title", "")
            st.session_state.f_publisher = ", ".join(p["name"] for p in data.get("publishers", []))
            places = data.get("publish_places", [])
            st.session_state.f_city = places[0]["name"] if places else ""
            st.session_state.lk_msg = ("Data ditemukan di Open Library. Nama penulis tertulis 'Nama Depan Nama Belakang', "
                                       "periksa dan perbaiki bila perlu.")
    except requests.HTTPError:
        st.session_state.lk_msg = "Data tidak ditemukan. Periksa DOI/ISBN atau isi manual."
    except requests.RequestException:
        st.session_state.lk_msg = "Tidak bisa terhubung ke internet. Isi manual saja."
    except Exception as ex:  # noqa: BLE001
        st.session_state.lk_msg = f"Terjadi masalah saat membaca data: {ex}"


def remove_entry(i):
    st.session_state.entries.pop(i)


# ───────── Dokumen Word ─────────
def add_runs(par, text):
    for chunk in re.split(r"(\*[^*]+\*)", text):
        if len(chunk) > 2 and chunk.startswith("*") and chunk.endswith("*"):
            par.add_run(chunk[1:-1]).italic = True
        elif chunk:
            par.add_run(chunk)


def build_docx(lines, style_name):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.add_run("DAFTAR PUSTAKA" if lang == "id" else "REFERENCES").bold = True
    for txt in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-1.27)
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        add_runs(p, txt)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ───────── Tab ─────────
tab_add, tab_list = st.tabs(["1. Tambah sumber", f"2. Daftar pustaka ({len(st.session_state.entries)})"])

with tab_add:
    with st.expander("Isi otomatis dari DOI atau ISBN (opsional)"):
        c1, c2 = st.columns([4, 1], vertical_alignment="bottom")
        c1.text_input("DOI atau ISBN", key="lk_id", placeholder="10.1000/xyz123  atau  9780471758013")
        c2.button("Cari data", on_click=do_lookup, width="stretch")
        if st.session_state.lk_msg:
            st.info(st.session_state.lk_msg)
        st.caption("Memakai Crossref dan Open Library (gratis, butuh internet). Hasil selalu perlu kamu periksa.")

    type_label = st.selectbox("Jenis sumber", list(TYPES.values()), key="f_type")
    kind = next(k for k, v in TYPES.items() if v == type_label)

    with st.form("form_sumber", clear_on_submit=True):
        a1, a2 = st.columns([3, 1])
        authors_txt = a1.text_area(
            "Penulis (satu per baris)",
            key="f_authors",
            height=110,
            placeholder="Hanafi, Ahmad Budi\nPutri, Sari\n\natau tanpa koma: Ahmad Budi Hanafi",
            help="Format 'Belakang, Depan' paling akurat. Tanpa koma, kata terakhir dianggap nama belakang.",
        )
        org = a2.checkbox("Penulis adalah lembaga", key="f_org", help="Contoh: Kementerian Kesehatan RI. Nama tidak akan disingkat.")

        b1, b2 = st.columns(2)
        year = b1.text_input("Tahun terbit", key="f_year", placeholder="2024")
        title = b2.text_input("Judul", key="f_title")

        e = {}
        if kind == "buku":
            c1, c2, c3 = st.columns(3)
            e["publisher"] = c1.text_input("Penerbit", key="f_publisher")
            e["city"] = c2.text_input("Kota terbit (untuk IEEE/Harvard)", key="f_city")
            e["edition"] = c3.text_input("Edisi (opsional)", key="f_edition", placeholder="2nd ed.")
            e["doi"] = ""
            e["url"] = st.text_input("URL (opsional)", key="f_url")
        elif kind == "jurnal":
            e["journal"] = st.text_input("Nama jurnal", key="f_journal")
            c1, c2, c3 = st.columns(3)
            e["volume"] = c1.text_input("Volume", key="f_volume")
            e["issue"] = c2.text_input("Nomor (issue)", key="f_issue")
            e["pages"] = c3.text_input("Halaman", key="f_pages", placeholder="45-58")
            c4, c5 = st.columns(2)
            e["doi"] = c4.text_input("DOI", key="f_doi", placeholder="10.1234/abcd")
            e["url"] = c5.text_input("URL (bila tanpa DOI)", key="f_url")
        elif kind == "website":
            c1, c2 = st.columns(2)
            e["site"] = c1.text_input("Nama situs / organisasi", key="f_site")
            e["url"] = c2.text_input("URL", key="f_url")
            c3, c4, c5 = st.columns(3)
            month = c3.selectbox("Bulan terbit (opsional)", range(13),
                                 format_func=lambda i: "-" if i == 0 else MONTHS[lang][i - 1], key="f_month")
            day = c4.number_input("Tanggal terbit (0 = kosong)", 0, 31, 0, key="f_day")
            acc = c5.date_input("Tanggal diakses", value=date.today(), key="f_acc")
            e["month"], e["day"], e["accessed"] = month, int(day), acc.isoformat()
        else:
            c1, c2, c3 = st.columns(3)
            e["degree"] = c1.selectbox("Jenis", ["Skripsi", "Tesis", "Disertasi"], key="f_degree")
            e["university"] = c2.text_input("Universitas", key="f_univ")
            e["city"] = c3.text_input("Kota (untuk IEEE)", key="f_city")
            e["url"] = st.text_input("URL repositori (opsional)", key="f_url")

        submitted = st.form_submit_button("Tambah ke daftar", type="primary")

    if submitted:
        entry = {"type": kind, "authors": parse_authors(authors_txt, org), "year": year.strip(),
                 "title": title.strip(), **{k: (v.strip() if isinstance(v, str) else v) for k, v in e.items()}}
        st.session_state.entries.append(entry)
        warns = check_entry(entry)
        st.success("Sumber ditambahkan. Lihat hasilnya di tab Daftar pustaka.")
        for w in warns:
            st.warning(w)

with tab_list:
    entries = st.session_state.entries
    with st.expander("Simpan / muat daftar sumber"):
        st.caption("Data hilang saat halaman ditutup. Simpan file ini agar bisa dilanjutkan lain waktu.")
        st.download_button("Simpan sumber (.json)", json.dumps(entries, ensure_ascii=False, indent=2),
                           "sumber_pustaka.json", "application/json", disabled=not entries)
        up = st.file_uploader("Muat file .json", type=["json"])
        if up is not None and st.button("Gabungkan ke daftar"):
            try:
                loaded = json.load(up)
                if not isinstance(loaded, list):
                    raise ValueError
                st.session_state.entries.extend(loaded)
                st.rerun()
            except (ValueError, json.JSONDecodeError):
                st.error("File tidak valid. Gunakan file .json yang disimpan dari aplikasi ini.")

    if not entries:
        st.info("Belum ada sumber. Tambahkan lewat tab Tambah sumber.")
    else:
        try:
            items = format_list(entries, style, lang)
        except Exception as ex:  # noqa: BLE001
            st.error(f"Ada sumber yang tidak bisa diformat: {ex}")
            items = []
        st.subheader(f"Gaya {style}")
        for idx, text in items:
            c1, c2 = st.columns([12, 1], vertical_alignment="center")
            c1.markdown(text.replace("$", r"\$"))
            c2.button("🗑", key=f"del{idx}", on_click=remove_entry, args=(idx,), help="Hapus sumber ini")

        lines = [t for _, t in items]
        plain = "\n\n".join(re.sub(r"\*", "", t) for t in lines)
        st.markdown("**Teks siap salin**")
        st.code(plain, language=None, wrap_lines=True)
        d1, d2, d3 = st.columns(3)
        d1.download_button("Unduh Word (.docx)", build_docx(lines, style), "daftar_pustaka.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        d2.download_button("Unduh teks (.txt)", plain, "daftar_pustaka.txt")
        if d3.button("Kosongkan semua"):
            st.session_state.entries = []
            st.rerun()
        st.caption("Di Word, teks miring sudah diterapkan. Pastikan hasilnya sesuai pedoman penulisan kampusmu.")
